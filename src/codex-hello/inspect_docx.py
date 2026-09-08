from docx import Document
from pathlib import Path

p = Path(r"C:\DEV\codex-hello\notice_kr.docx")
doc = Document(p)
for i, para in enumerate(doc.paragraphs):
    print(f"P{i} style={para.style.name!r} numPr={para._p.pPr.numPr is not None if para._p.pPr is not None else False} text={para.text!r}")
    for j, run in enumerate(para.runs):
        print(f"  R{j} bold={run.bold} italic={run.italic} underline={run.underline} text={run.text!r}")
for ti, table in enumerate(doc.tables):
    print(f"TABLE {ti}")
    for ri, row in enumerate(table.rows):
        print(" | ".join(cell.text.replace("\n", " / ") for cell in row.cells))
for si, section in enumerate(doc.sections):
    print(f"HEADER {si}: {[p.text for p in section.header.paragraphs]}")
    print(f"FOOTER {si}: {[p.text for p in section.footer.paragraphs]}")
